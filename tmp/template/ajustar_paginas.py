from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
d=Document('RelatórioT0_Final.docx')
for p in list(d.paragraphs):
 t=p.text.strip()
 if t in ['RESUMO','LISTA DE FIGURAS','LISTA DE TABELAS','SUMÁRIO']:
  prev=p._p.getprevious()
  while prev is not None and prev.tag==qn('w:p') and not ''.join(prev.itertext()).strip() and not prev.xpath('./w:pPr/w:sectPr'):
   old=prev;prev=prev.getprevious();old.getparent().remove(old)
  # Remove preceding empty break paragraphs to avoid double page breaks.
  if prev is not None and prev.tag==qn('w:p') and not prev.xpath('.//w:t') and not prev.xpath('./w:pPr/w:sectPr'):
   prev.getparent().remove(prev)
  p.paragraph_format.page_break_before=False;p.paragraph_format.keep_with_next=True
  br=p.insert_paragraph_before('');br.paragraph_format.keep_with_next=False;br.paragraph_format.space_before=Pt(0);br.paragraph_format.space_after=Pt(0);br.paragraph_format.line_spacing=1;br.add_run().add_break(WD_BREAK.PAGE)
 if t=='2026':p.text='2026';p.paragraph_format.keep_with_next=False
 if t.startswith('Fonte: Acervo fotográfico'):p.text='Fonte: Acervo fotográfico do projeto (15 set. 2026).'
 if t.startswith('ACERVO fotográfico'):p.text='ACERVO fotográfico do projeto de Física Experimental I. Uberlândia, 15 set. 2026. 3 fotografias digitais.'
 if t.startswith('As medidas foram expressas em centímetros.'):
  p.text='As medidas foram expressas em centímetros. Adotou-se a contribuição instrumental de 0,01 milímetro, equivalente a 0,001 centímetro, para a bateria, e de 0,05 milímetro, equivalente a 0,005 centímetro, para a arruela. Para as duas peças com ressalto, adotou-se o valor de 0,005 centímetro, pois esses itens não especificam outro valor.'
 if t.startswith('A propagação empregada é uma aproximação'):
  p.text=t.replace('A contribuição instrumental adotada para as peças compostas permanece uma hipótese do tratamento.','Para as peças compostas, utilizou-se a contribuição instrumental adotada de 0,005 centímetro.')
 if t.upper()=='1 INTRODUÇÃO' and p.style.name.startswith('Heading'):p.text='1 INTRODUÇÃO'
d.save('RelatórioT0_Final.docx')
