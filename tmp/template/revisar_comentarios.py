from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
import re
src=Path('RelatórioT0_Template.docx');d=Document(src)
# Remove the abstract and inline editing notes between keywords and figure list.
ps=d.paragraphs;a=next(j for j,p in enumerate(ps) if p.text.startswith('Palavras-chave:'));b=next(j for j,p in enumerate(ps) if p.text=='LISTA DE FIGURAS')
for p in ps[a+1:b]:p._p.getparent().remove(p._p)
d.paragraphs[a+1].paragraph_format.page_break_before=True
for p in list(d.paragraphs):
 t=p.text
 if any(t.strip().lstrip('.… ').startswith(x) for x in ['Tire essa parte','Atravessa, voce','Escreva em números','Pode apagar isso']):
  p._p.getparent().remove(p._p);continue
 new=t
 if t.startswith('Este trabalho determina'):
  new=t.replace('aproximados de um, dezenove centésimos, onze inteiros e sete décimos e três centímetros cúbicos','aproximados de 1,00; 0,19; 11,7 e 3,0 cm³').replace('de dois inteiros e cinquenta e seis centésimos por cento a quarenta e seis inteiros e cinquenta e cinco centésimos por cento','de 2,56% a 46,55%').replace('Os resultados das peças compostas permanecem condicionados à hipótese geométrica de furo com diâmetro uniforme ao longo de toda a altura considerada.','As peças compostas possuem furo passante, representado no modelo por um diâmetro uniforme ao longo da altura total.')
 if t.startswith('O tratamento seguiu cinco etapas:'):
  new=t.replace('Todas as leituras confirmadas foram mantidas, incluindo a segunda medida de L igual a 1,31 centímetro. ','')
 if t.startswith('As peças compostas foram modeladas'):
  new='As peças compostas possuem furo passante, que atravessa toda a altura. Por isso, foram modeladas como duas arruelas de mesmo diâmetro interno, representando a base e o ressalto.'
 if t.startswith('As Figuras 1 a 3 apresentam'):new='As Figuras 1 a 3 apresentam os materiais e dois registros de posicionamento do paquímetro.'
 if t.startswith('A Tabela 2 apresenta os resultados'):new='A Tabela 2 apresenta os resultados das quatro primeiras etapas. As casas decimais adicionais permitem acompanhar os cálculos intermediários.'
 if t.startswith('Para a peça com ressalto, a arruela superior'):new=t.replace('Sob a hipótese descrita na metodologia, o volume total corresponde à soma:','O volume total corresponde à soma:')
 if t.startswith('Na peça com corte plano, subtrai-se'):
  new='Na peça com corte plano, subtrai-se o volume de um segmento circular da base. L representa a linha reta que une as extremidades do corte na circunferência, chamada corda. Nas equações a seguir, o ângulo θ é expresso em radianos.'
 if t.startswith('A bateria apresentou a menor incerteza relativa'):
  new=t.split('Esses valores')[0].strip()+' Essas incertezas avaliam a precisão dos resultados; a comparação com o volume verdadeiro exigiria um valor de referência.'
 if t.startswith('A propagação empregada é uma aproximação'):
  new=t.replace('Além dessas limitações, os resultados das peças compostas dependem da confirmação da extensão do furo e da contribuição instrumental adotada.','A contribuição instrumental adotada para as peças compostas permanece uma hipótese do tratamento.')
 if t.startswith('A dispersão das alturas foi a principal limitação'):
  new=t.replace('Antes de interpretar como definitivos os resultados das peças compostas, é necessário confirmar se o furo atravessa a altura considerada no modelo. ','')
 if t.startswith('ACERVO fotográfico'):new='ACERVO fotográfico do projeto de Física Experimental I. [S. l.: s. n.], [s. d.]. 3 fotografias digitais.'
 if t.startswith(('IWAMOTO,','ODASHIMA,','SOUSA,','TEMPLATE de trabalhos')):new=t.split(' Arquivo:')[0]
 if new!=t:p.text=new
# Remove residual editorial blank paragraphs, preserving page/section breaks and pictures.
for p in list(d.paragraphs):
 if not p.text.strip() and not p._p.xpath('.//w:br | .//w:sectPr | .//w:drawing | .//m:oMath | .//w:fldChar'):
  # retain deliberate front matter spacing
  if p.paragraph_format.space_after is None or p.paragraph_format.space_after.pt<20:
   p._p.getparent().remove(p._p)
for el in list(d.element.xpath('//w:highlight')):el.getparent().remove(el)
# Save separately because the user's annotated document is currently open in Word.
out=Path('RelatórioT0_Revisado.docx');d.save(out)
assert len(d.element.xpath('//m:oMath'))==17
assert len(d.inline_shapes)==3
text='\n'.join(p.text for p in d.paragraphs)
for banned in ['ABSTRACT','Pode apagar isso','Pra que isso','Isso é necessario','preciso confirmar','confirmar se o furo','Arquivo:','refs/Fotos']:
 assert banned not in text,banned
print(out)
