from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH as A,WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy
p=Path('RelatórioT0_Template.docx');d=Document(p)
def insert_before(anchor,text='',style='Normal'):
 p=anchor.insert_paragraph_before(text,style);return p
def plain(p):
 p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
 for r in p.runs:r.font.size=Pt(10)
def field(p,code):
 for kind in ['begin','instr','separate','end']:
  r=OxmlElement('w:r')
  if kind=='instr':el=OxmlElement('w:instrText');el.text=' '+code+' ';el.set(qn('xml:space'),'preserve')
  else:el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),kind)
  r.append(el);p._p.append(r)
def bookmark(p,name,k):
 s=OxmlElement('w:bookmarkStart');s.set(qn('w:id'),str(k));s.set(qn('w:name'),name);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(k));p._p.insert(0,s);p._p.append(e)
# Add source citations to the existing passages, retaining paragraph styles.
for par in d.paragraphs:
 if par.text.startswith('Para três leituras de uma dimensão'):
  par.text=par.text.replace('(Iwamoto et al., 2014).','(Iwamoto et al., 2014; Odashima, [s. d.]b).')
 if par.text.startswith('Nessa expressão, m é'):
  par.text=par.text.rstrip('.')+' (Odashima, [s. d.]a).'
 if par.text.startswith('A análise também compara'):
  new=OxmlElement('w:p');par._p.addnext(new)
  from docx.text.paragraph import Paragraph
  np=Paragraph(new,par._parent);np.style=d.styles['Normal'];np.text='A organização da apresentação dos resultados segue as orientações de elaboração de relatórios da disciplina (Odashima, [s. d.]c). A formatação utiliza o template de trabalhos acadêmicos disponibilizado entre os materiais do projeto (Template..., [s. d.]).'
# Figure list before table list. Move the existing preceding page break to our new list, and add another for tables.
anchor=next(x for x in d.paragraphs if x.text=='LISTA DE TABELAS')
p=insert_before(anchor,'LISTA DE FIGURAS','Resumo');p.alignment=A.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(18)
for r in p.runs:r.bold=True
figures=[('Instrumentos e peças sobre a bancada','Foto_objetos_de_medidas.jpeg','A fotografia apresenta as quatro peças, o paquímetro, a régua e a trena sobre a bancada. A presença dos instrumentos documenta a organização do material; os cálculos deste relatório utilizam as leituras de paquímetro registradas no roteiro.'),('Posicionamento do paquímetro para medição externa','Foto_medicao01.jpeg','O registro mostra o posicionamento das faces de medição externa do paquímetro junto à peça com ressalto e furo.'),('Posicionamento da extremidade do paquímetro junto à peça','Foto_medicao02.jpeg','A fotografia mostra a extremidade do paquímetro posicionada junto à peça com corte plano. A imagem documenta o procedimento, sem fornecer uma leitura numérica suficientemente clara para substituir os valores da tabela de medidas.')]
for j,(title,_,_) in enumerate(figures,1):
 p=insert_before(anchor,f'Figura {j} – {title}\t');p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT);field(p,f'PAGEREF fig{j} \\h')
p=insert_before(anchor);p.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
# Insert photographs immediately before the page break preceding measurement table.
tab=next(x for x in d.paragraphs if x.text.startswith('Tabela 1 – Medidas') and '\t' not in x.text)
prev=tab._p.getprevious()
from docx.text.paragraph import Paragraph
anchor=Paragraph(prev,tab._parent) if prev.tag==qn('w:p') else tab
for j,(title,file,description) in enumerate(figures,1):
 p=insert_before(anchor);p.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
 if j==1:
  p=insert_before(anchor,'3.1 Registros fotográficos','Heading 2')
  insert_before(anchor,'As Figuras 1 a 3 apresentam os materiais e dois registros de posicionamento do paquímetro. As imagens integram o acervo fotográfico disponibilizado com os dados do projeto (Acervo..., [s. d.]).')
 p=insert_before(anchor,f'Figura {j} – {title}');plain(p);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(6);bookmark(p,'fig'+str(j),400+j)
 p=insert_before(anchor);p.alignment=A.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
 p.add_run().add_picture(str(Path('refs/Fotos')/file),width=Cm(10.5 if j==1 else 12))
 p=insert_before(anchor,'Fonte: Acervo fotográfico do projeto ([s. d.]).');plain(p);p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_with_next=True
 insert_before(anchor,description)
# Replace only bibliography paragraphs after the existing reference heading.
ref=next(x for x in d.paragraphs if x.text=='REFERÊNCIAS' and x.style.name=='Seção não numerada')
node=ref._p.getnext()
while node is not None and node.tag!=qn('w:sectPr'):
 nxt=node.getnext();node.getparent().remove(node);node=nxt
refs=[
'ACERVO fotográfico do projeto de Física Experimental I. [S. l.: s. n.], [s. d.]. 3 fotografias digitais: Foto_objetos_de_medidas.jpeg, Foto_medicao01.jpeg e Foto_medicao02.jpeg. Arquivos disponibilizados na pasta refs/Fotos.',
'IWAMOTO, Wellington Akira; GUARANY, Cristiano Alves; FOSCHINI, Mauricio; DI LORENZO, Antonino. Guias e roteiros para Laboratório de Física Experimental I. 1. ed. Uberlândia: Instituto de Física, Universidade Federal de Uberlândia, 2014. Arquivo: Apostila_LAB1.pdf.',
'ODASHIMA, Mariana M. Física Experimental 1: Laboratório de Física: Mecânica. Aula 1: apresentação do curso, teoria de erros e medidas, algarismos significativos. Uberlândia: Instituto de Física, Universidade Federal de Uberlândia, [s. d.]a. 32 slides. Arquivo: Aula1-AlgSignificativos.pdf.',
'ODASHIMA, Mariana M. Física Experimental 1: Lab Física: Mecânica. Aula 2: teoria de erros; erro instrumental, estatístico, total. Uberlândia: Instituto de Física, Universidade Federal de Uberlândia, [s. d.]b. 25 slides. Arquivo: Aula2-ErrosEstatistica-simpl.pdf.',
'ODASHIMA, Mariana M. Física Geral Experimental: Lab Física 1: Mecânica. Aula 4: como fazer relatório. Uberlândia: Instituto de Física, Universidade Federal de Uberlândia, [s. d.]c. 25 slides. Arquivo: Aula4-ComoCriarRelatorios.pdf.',
'SOUSA, Lucas Soares. L1-Laboratório-FG1: Engenharia – Diurno. Roteiro de atividades. 15 set. 2026. 5 p. Arquivo: Trabalho_FG1.pdf.',
'TEMPLATE de trabalhos acadêmicos. [S. l.: s. n.], [s. d.]. Modelo de formatação em arquivo Word. Arquivo: template_tcc.docx.'
]
for text in refs:
 p=d.add_paragraph(text);p.alignment=A.LEFT;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(12)
d.save('RelatórioT0_Template.docx')
print('Added 3 photographs, figure list and 7 reference entries.')
