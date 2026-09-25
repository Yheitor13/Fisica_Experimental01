from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from pathlib import Path
names='Arthur Borges Fernandes\nEnzo Cerezine Oliveira\nHeitor Yochida de Ávila\nSofia Skolimoski\nTiago de Almeida Raile'
d=Document('RelatórioT0_Revisado.docx');ps=d.paragraphs
ps[3].text=names
ps[18].text=names;ps[18].alignment=A.CENTER;ps[18].paragraph_format.space_after=Pt(0);ps[18].paragraph_format.line_spacing=1.5;ps[18].paragraph_format.first_line_indent=Cm(0)
for p in [ps[3],ps[18]]:
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
for idx,space in [(2,30),(4,40),(8,160),(19,32),(23,24),(30,36)]:ps[idx].paragraph_format.space_after=Pt(space)
for idx in [9,10,11,12,13,14,31,32,33,34]:ps[idx]._p.getparent().remove(ps[idx]._p)
replacements={'Curso: [.......]':'Curso: Engenharia Eletrônica e de Telecomunicações / Engenharia de Controle e Automação','Turma: [.......]':'Turma: LabFis1_Eng_2026-2','Docente: [.......]':'Docente: Lucas Soares Sousa','Data de realização: [.......]':'Data de realização: 15 de setembro de 2026'}
for p in d.paragraphs:
 if p.text in replacements:p.text=replacements[p.text]
 if p.text=='RESUMO':p.paragraph_format.page_break_before=True
 if p.text.startswith('Palavras-chave:'):p.paragraph_format.space_before=Pt(18)
d.core_properties.author='Arthur Borges Fernandes; Enzo Cerezine Oliveira; Heitor Yochida de Ávila; Sofia Skolimoski; Tiago de Almeida Raile'
d.save('RelatórioT0_Final.docx')
print('Equations',len(d.element.xpath('//m:oMath')),'Photos',len(d.inline_shapes))
print('Pending:',[p.text for p in d.paragraphs if '[.......]' in p.text])
